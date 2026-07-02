"""Export a session's AGENT TELEMETRY (tools · MCP servers · tokens) to Excel, labelled with
session / date / time / topic. A BASELINE snapshot to compare before vs after a change (e.g. adding the
Planner agent + scoping the Maker's authority). READ-ONLY -- reads the persisted activity, writes a
spreadsheet. Falls back to CSV if openpyxl isn't installed.

    python eval_telemetry.py <session_id>        ->  evals/telemetry_<session>.xlsx
"""
import sys
import json
import csv
import collections
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()
from session import SESSIONS_DIR
from session_title import cached_title
import session_meta

ROLE = {"doer": "Maker (+ Heal)", "verifier": "Verifier", "conductor": "Conductor",
        "planner": "Planner", "engineering": "Board — Engineering", "procurement": "Board — Procurement",
        "compliance": "Board — Compliance", "finance": "Board — Finance",
        "planning": "Board — Planning", "chair": "Chair", "guide": "Guide", "maker": "Maker", "heal": "Heal"}
ORDER = ["doer", "maker", "heal", "planner", "verifier", "conductor",
         "engineering", "procurement", "compliance", "finance", "planning", "chair", "guide"]


def aggregate(sid):
    d = SESSIONS_DIR / sid
    ap = d / "activity.jsonl"
    if not ap.exists():
        raise SystemExit(f"no activity for session {sid}")
    turns = [json.loads(l) for l in ap.read_text(encoding="utf-8").splitlines() if l.strip()]
    tools = collections.defaultdict(collections.Counter)
    mcp = collections.defaultdict(collections.Counter)
    tok = collections.defaultdict(lambda: [0, 0, 0])
    for t in turns:
        for e in t.get("chain", []):
            a, k = e.get("agent"), e.get("kind")
            if k == "tool":
                tools[a][e.get("tool")] += 1
            elif k == "usage":
                tok[a][0] += e.get("tokens_in", 0); tok[a][1] += e.get("tokens_out", 0); tok[a][2] += e.get("calls", 0)
                for s, n in (e.get("mcp") or {}).items():
                    mcp[a][s] += n
    # date / time / topic
    c_iso, u_iso = session_meta.times(d)
    if not u_iso:
        import datetime
        u_iso = datetime.datetime.fromtimestamp((d / "trace.jsonl").stat().st_mtime).isoformat(timespec="seconds")
    first = ""
    tp = d / "trace.jsonl"
    if tp.exists():
        for l in tp.read_text(encoding="utf-8").splitlines():
            try:
                r = json.loads(l)
            except Exception:
                continue
            if r.get("role") == "user" and r.get("text"):
                first = r["text"].strip(); break
    date, time = (u_iso.split("T") + [""])[:2]
    meta = {"session": sid, "date": date, "time": time, "topic": cached_title(d) or first[:60],
            "turns": len(turns), "user_prompt": first[:80]}
    # per-agent rows
    agents = [a for a in ORDER if a in set(list(tools) + list(tok))]
    agents += [a for a in (set(list(tools) + list(tok)) - set(agents)) if a]
    rows = []
    for a in agents:
        ti, to, tc = tok[a]
        rows.append({"agent": a, "role": ROLE.get(a, a),
                     "tool_calls": sum(tools[a].values()),
                     "distinct_tools": len(tools[a]),
                     "tools": ", ".join(f"{k}×{v}" for k, v in tools[a].most_common()),
                     "mcp_servers": ", ".join(f"{k}×{v}" for k, v in mcp[a].most_common()),
                     "tokens_in": ti, "tokens_out": to, "model_calls": tc})
    return meta, rows


def write_xlsx(path, meta, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active; ws.title = "Telemetry"
    hdrfill = PatternFill("solid", fgColor="1A3C5E"); hdrfont = Font(color="FFFFFF", bold=True)
    r = 1
    ws.cell(r, 1, "SESSION TELEMETRY BASELINE").font = Font(bold=True, size=13); r += 1
    for k in ("session", "date", "time", "topic", "turns", "user_prompt"):
        ws.cell(r, 1, k).font = Font(bold=True); ws.cell(r, 2, str(meta[k])); r += 1
    r += 1
    cols = ["agent", "role", "tool_calls", "distinct_tools", "tools", "mcp_servers",
            "tokens_in", "tokens_out", "model_calls"]
    for c, name in enumerate(cols, 1):
        cell = ws.cell(r, c, name); cell.fill = hdrfill; cell.font = hdrfont
    r += 1
    for row in rows:
        for c, name in enumerate(cols, 1):
            ws.cell(r, c, row[name])
        r += 1
    # totals
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws.cell(r, 7, sum(x["tokens_in"] for x in rows)).font = Font(bold=True)
    ws.cell(r, 8, sum(x["tokens_out"] for x in rows)).font = Font(bold=True)
    ws.cell(r, 3, sum(x["tool_calls"] for x in rows)).font = Font(bold=True)
    widths = [14, 22, 11, 14, 60, 46, 11, 11, 11]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)


def write_csv(path, meta, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for k in ("session", "date", "time", "topic", "turns", "user_prompt"):
            w.writerow([k, meta[k]])
        w.writerow([])
        cols = ["agent", "role", "tool_calls", "distinct_tools", "tools", "mcp_servers",
                "tokens_in", "tokens_out", "model_calls"]
        w.writerow(cols)
        for row in rows:
            w.writerow([row[c] for c in cols])


def _persist_s3(sid, meta, rows, local_file):
    """Best-effort: push the export (xlsx/csv) + a JSON of the data to S3 d2m/evals/telemetry_<sid>/ so
    the baseline is durable and comparable later. No-op if no bucket/creds."""
    import os
    if not os.getenv("S3_BUCKET"):
        return
    try:
        import boto3
        c = boto3.client("s3", region_name=os.getenv("S3_REGION", "us-west-1"))
        bk = os.getenv("S3_BUCKET"); px = os.getenv("S3_PREFIX", "d2m")
        base = f"{px}/evals/telemetry_{sid}"
        c.put_object(Bucket=bk, Key=f"{base}/telemetry.json",
                     Body=json.dumps({"meta": meta, "agents": rows}, indent=2).encode())
        with open(local_file, "rb") as f:
            c.put_object(Bucket=bk, Key=f"{base}/{local_file.name}", Body=f.read())
        print(f"  -> s3://{bk}/{base}/ (json + {local_file.suffix[1:]})")
    except Exception as e:
        print(f"  (S3 persist skipped: {e})")


def main():
    if len(sys.argv) < 2:
        print("usage: python eval_telemetry.py <session_id>"); return
    sid = sys.argv[1]
    meta, rows = aggregate(sid)
    out = Path("evals"); out.mkdir(exist_ok=True)
    print(f"session {sid} · {meta['date']} {meta['time']} · {meta['topic']}")
    print(f"  {len(rows)} agents · {sum(r['tokens_in'] for r in rows):,} in + "
          f"{sum(r['tokens_out'] for r in rows):,} out tokens")
    try:
        p = out / f"telemetry_{sid}.xlsx"
        write_xlsx(p, meta, rows)
    except ImportError:
        p = out / f"telemetry_{sid}.csv"
        write_csv(p, meta, rows)
        print("  (openpyxl not installed -> CSV; opens in Excel)")
    print(f"  -> {p}")
    _persist_s3(sid, meta, rows, p)


if __name__ == "__main__":
    main()
