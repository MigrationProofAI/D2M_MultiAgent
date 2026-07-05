#!/usr/bin/env python
"""make_altivar_bom.py -- a REALISTIC 3-level Schneider Electric product for the multi-level showcase:
an Altivar variable-frequency drive. FERT -> L1 modules (HALB) -> L2 sub-assemblies (HALB) -> L3 parts.
Every made node (FERT + L1 + L2) gets its OWN BOM/routing/PV via the recursive genesis. Vendors default
to the one known-real supplier (17300001). Same two-sheet schema as the other fixtures.

  python make_altivar_bom.py --out bom_altivar.xlsx
"""
import os
import sys
import argparse

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "mcp_server"))
BOM_COLS = ["id", "parent_id", "level", "role", "type", "name", "description",
            "material", "quantity", "unit", "vendor", "price", "plant"]
OP_COLS = ["node_id", "operation", "text", "work_center", "setup_time", "run_time"]

FERT = ("Altivar ATV630 22kW VFD", "Altivar ATV630 400V 22kW drive",
        [("10", "Module Marriage", "ASSEMBLY", 40, 30), ("20", "Wiring & Loom", "ASSEMBLY", 20, 15),
         ("30", "Functional & HiPot Test", "PACK01", 25, 20), ("40", "Pack & Crate", "PACK01", 6, 8)])

# L1 module = (name, desc, ops, [L2 made sub-assemblies], [L2 direct-bought parts])
#   L2 sub = (name, desc, ops, [L3 bought parts])
#   bought part = (name, desc, type, qty, price)
_ASM = [("10", "Assemble", "ASSEMBLY", 15, 8), ("20", "Test", "PACK01", 8, 5)]
MODULES = [
 ("Power Module", "400V 22kW power stage", _ASM,
  [("Rectifier Stage", "3-phase input rectifier", _ASM,
    [("Diode Bridge Module", "3-ph 1200V diode bridge", "HAWA", 1, 21.0),
     ("Input Snubber Cap", "630V film snubber cap", "HAWA", 3, 2.4),
     ("Inrush Thermistor", "NTC inrush limiter", "HAWA", 1, 1.8),
     ("Rectifier Bus Bar", "tin-plated copper bus bar", "ROH", 1, 6.5),
     ("MOV Surge Set", "metal-oxide varistor set", "HAWA", 1, 3.2)]),
   ("Inverter Stage", "IGBT inverter stage", _ASM,
    [("IGBT Module", "1200V 75A IGBT module", "HAWA", 6, 24.5),
     ("Gate Driver Board", "6-ch isolated gate driver", "HAWA", 1, 34.0),
     ("Output Bus Bar", "laminated output bus bar", "ROH", 1, 9.0),
     ("Current Sensor", "hall-effect current sensor", "HAWA", 3, 6.2),
     ("Gate Resistor Set", "gate + DESAT resistor set", "ROH", 1, 1.9),
     ("Snubber Cap Set", "IGBT snubber cap set", "HAWA", 1, 4.1)]),
   ("DC Bus Assembly", "DC-link capacitor bank", _ASM,
    [("DC-Link Capacitor", "450V 2200uF electrolytic", "HAWA", 4, 8.6),
     ("Balancing Resistor", "22k 5W balancing resistor", "ROH", 4, 0.7),
     ("Pre-charge Resistor", "50R 50W pre-charge resistor", "HAWA", 1, 3.4),
     ("DC Bus Bar", "laminated DC bus bar", "ROH", 1, 8.0)])],
  [("Power Base Plate", "AlSiC power base plate", "ROH", 1, 12.0),
   ("Thermal Interface Pad", "graphite thermal pad", "HAWA", 2, 1.6)]),
 ("Control Module", "digital control unit", _ASM,
  [("Control PCBA", "main control board", _ASM,
    [("Control MCU", "32-bit motor-control MCU", "HAWA", 1, 12.5),
     ("FPGA", "PWM/logic FPGA", "HAWA", 1, 18.0),
     ("DDR Memory", "DDR3 512MB", "HAWA", 1, 6.0),
     ("Flash Memory", "QSPI 128Mb flash", "HAWA", 1, 2.8),
     ("Crystal Oscillator", "25MHz TCXO", "HAWA", 1, 1.4),
     ("Control Connector Set", "board-to-board connector set", "HAWA", 1, 3.6)]),
   ("Power Supply Board", "24/15/5V SMPS", _ASM,
    [("SMPS Transformer", "flyback transformer", "HAWA", 1, 5.5),
     ("Buck Regulator", "3A synchronous buck IC", "HAWA", 2, 2.2),
     ("Bulk Cap Set", "low-ESR bulk cap set", "HAWA", 1, 3.0),
     ("Rectifier Diode Set", "schottky rectifier set", "HAWA", 1, 1.7),
     ("Opto-coupler", "feedback opto-coupler", "HAWA", 1, 0.9)]),
   ("I/O Board", "digital + analog I/O", _ASM,
    [("Signal Relay", "24V signal relay", "HAWA", 3, 1.5),
     ("Opto-isolator", "digital-input opto-isolator", "HAWA", 4, 0.8),
     ("Analog Input Terminal", "0-10V/4-20mA input terminal", "HAWA", 1, 2.1),
     ("Digital I/O Terminal", "24V digital I/O terminal", "HAWA", 1, 2.3)])],
  [("Control Backplane", "control backplane PCB", "ROH", 1, 7.0),
   ("EMC Shield Can", "stamped EMC shield", "ROH", 1, 2.5)]),
 ("Cooling Assembly", "forced-air cooling", _ASM,
  [("Heatsink Sub", "extruded aluminium heatsink", _ASM,
    [("Extruded Heatsink", "black-anodised extruded heatsink", "ROH", 1, 22.0),
     ("Thermal Grease", "high-k thermal grease", "HAWA", 1, 1.2),
     ("Heatsink Clip Set", "spring clip set", "ROH", 1, 1.1)]),
   ("Fan Unit", "dual axial fan", _ASM,
    [("Axial Fan", "24V 120mm axial fan", "HAWA", 2, 9.5),
     ("Fan Guard", "steel fan guard", "ROH", 2, 0.9),
     ("Fan Wiring Harness", "fan power harness", "ROH", 1, 2.0),
     ("Fan Connector", "3-pin fan connector", "HAWA", 2, 0.6)])],
  [("Air Duct", "moulded air duct", "ROH", 1, 4.5),
   ("Dust Filter", "IP54 dust filter", "HAWA", 1, 2.0)]),
 ("Enclosure Assembly", "IP20 enclosure", _ASM,
  [("Front Cover Sub", "front cover + gasket", _ASM,
    [("Front Cover Panel", "ABS front cover", "ROH", 1, 6.8),
     ("Cover Gasket", "EPDM cover gasket", "HAWA", 1, 1.3),
     ("Cover Screw Set", "captive screw set", "ROH", 1, 0.8)])],
  [("Chassis Frame", "galvanised chassis frame", "ROH", 1, 14.0),
   ("Rear Panel", "DIN/wall mount rear panel", "ROH", 1, 5.5),
   ("Side Panel", "vented side panel", "ROH", 2, 3.2),
   ("DIN Rail Clip Set", "35mm DIN clip set", "HAWA", 1, 1.5),
   ("Grounding Stud", "M6 grounding stud", "HAWA", 1, 0.7),
   ("Nameplate Label", "laser-etched nameplate", "HAWA", 1, 1.0)]),
 ("HMI Module", "graphic display terminal", _ASM,
  [("Display Board", "LCD + keypad board", _ASM,
    [("Graphic LCD", "240x160 graphic LCD", "HAWA", 1, 18.0),
     ("Keypad Membrane", "tactile keypad membrane", "HAWA", 1, 3.5),
     ("Display Driver IC", "LCD driver IC", "HAWA", 1, 2.6),
     ("Backlight LED Set", "white backlight LED set", "HAWA", 1, 1.4),
     ("Display Connector", "FFC display connector", "HAWA", 1, 0.9)])],
  [("HMI Bezel", "moulded HMI bezel", "ROH", 1, 3.0),
   ("HMI Cable", "RJ45 HMI cable", "HAWA", 1, 2.4),
   ("HMI Mounting Clip", "panel mount clip set", "HAWA", 1, 1.1)]),
 ("Braking Unit", "dynamic braking chopper", _ASM,
  [("Brake Chopper Sub", "brake IGBT chopper", _ASM,
    [("Brake IGBT", "1200V 50A brake IGBT", "HAWA", 1, 16.0),
     ("Brake Diode", "1200V freewheel diode", "HAWA", 1, 4.5),
     ("Brake Snubber Cap", "brake snubber cap", "HAWA", 1, 2.0),
     ("Brake Gate Board", "single-ch gate driver", "HAWA", 1, 8.0)])],
  [("Brake Terminal", "brake resistor terminal", "HAWA", 1, 2.2),
   ("Brake Bus Link", "brake bus link bar", "ROH", 1, 2.8)]),
 ("Terminal Module", "power + control terminals", _ASM,
  [("Power Terminal Block", "input/output power terminals", _ASM,
    [("Power Terminal Lug", "M8 power terminal lug", "HAWA", 6, 1.9),
     ("Terminal Insulator", "moulded terminal insulator", "ROH", 1, 2.4),
     ("Terminal Screw Set", "M8 terminal screw set", "ROH", 1, 1.0)])],
  [("Control Terminal Block", "pluggable control terminal", "HAWA", 1, 4.0),
   ("Ferrule Kit", "wire ferrule kit", "HAWA", 1, 1.6),
   ("Cable Gland Set", "IP54 cable gland set", "HAWA", 1, 2.8),
   ("Terminal Cover", "clear terminal cover", "ROH", 1, 1.5)]),
]
TOP_BOUGHT = [("Fastener Kit", "drive-level fastener kit", "ROH", 1, 3.5),
              ("Owner Manual", "printed manual + safety guide", "HAWA", 1, 1.5),
              ("Serial Plate", "laser serial nameplate", "ROH", 1, 1.2)]


def build(vendor, plant):
    bom, ops = [], []

    def row(**kw):
        bom.append({c: kw.get(c, "") for c in BOM_COLS})

    def op(nid, o):
        for (num, text, wc, su, ru) in o:
            ops.append({"node_id": nid, "operation": num, "text": text,
                        "work_center": wc, "setup_time": su, "run_time": ru})

    row(id="P", parent_id="", level=0, role="parent", type="FERT",
        name=FERT[0], description=FERT[1], quantity=1, unit="EA", plant=plant)
    op("P", FERT[2])
    for i, (mname, mdesc, mops, subs, bought) in enumerate(MODULES):
        mid = f"M{i}"
        row(id=mid, parent_id="P", level=1, role="made", type="HALB", name=mname, description=mdesc, quantity=1, unit="EA")
        op(mid, mops)
        for j, (sname, sdesc, sops, parts) in enumerate(subs):
            sid = f"{mid}S{j}"
            row(id=sid, parent_id=mid, level=2, role="made", type="HALB", name=sname, description=sdesc, quantity=1, unit="EA")
            op(sid, sops)
            for k, (pn, pd, pt, qty, pr) in enumerate(parts):
                row(id=f"{sid}P{k}", parent_id=sid, level=3, role="bought", type=pt,
                    name=pn, description=pd, quantity=qty, unit="EA", vendor=vendor, price=pr)
        for k, (pn, pd, pt, qty, pr) in enumerate(bought):
            row(id=f"{mid}B{k}", parent_id=mid, level=2, role="bought", type=pt,
                name=pn, description=pd, quantity=qty, unit="EA", vendor=vendor, price=pr)
    for k, (pn, pd, pt, qty, pr) in enumerate(TOP_BOUGHT):
        row(id=f"B{k}", parent_id="P", level=1, role="bought", type=pt,
            name=pn, description=pd, quantity=qty, unit="EA", vendor=vendor, price=pr)
    return bom, ops


def write_xlsx(path, bom_rows, op_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "BOM"; ws.append(BOM_COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in bom_rows:
        ws.append([r.get(col, "") for col in BOM_COLS])
    wo = wb.create_sheet("Operations"); wo.append(OP_COLS)
    for c in wo[1]:
        c.font = Font(bold=True)
    for r in op_rows:
        wo.append([r.get(col, "") for col in OP_COLS])
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="bom_altivar.xlsx")
    ap.add_argument("--vendor", default="17300001")
    ap.add_argument("--plant", default=os.getenv("SAP_PLANT", "1710"))
    args = ap.parse_args()
    bom, ops = build(args.vendor, args.plant)
    write_xlsx(args.out, bom, ops)
    from excel_bom import genesis_from_excel
    spec, warnings = genesis_from_excel(args.out)

    def count(cs, made=0, bought=0, depth=1, maxd=1):
        for c in cs:
            maxd = max(maxd, depth)
            if c.get("role") == "made" or c.get("type") in ("HALB", "FERT"):
                made += 1
            else:
                bought += 1
            made, bought, maxd = count(c.get("components", []), made, bought, depth + 1, maxd)
        return made, bought, maxd
    m, b, d = count(spec["components"])
    print(f"wrote {args.out}: {len(bom)} rows | 1 FERT + {m} HALB + {b} bought = {1 + m + b} materials, max depth {d + 1}")
    for w in warnings[:6]:
        print(f"  WARN: {w}")


if __name__ == "__main__":
    main()
