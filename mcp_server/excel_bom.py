"""excel_bom.py -- deterministic tabbed-Excel -> genesis spec (the structured alternative to the vision
model). A workbook with a `BOM` sheet (the hierarchy, one row per node) + an optional `Operations` sheet
(routing per made node) is parsed straight into the SAME multi-level spec that run_genesis consumes -- no
vision, no LLM, exact + reproducible. This is the Excel analogue of genesis_from_csv (genesis.py), extended
to a real tree + operations, so genesis can be scale-tested at 50/100/150/200 materials. Read-only: NO SAP
calls -- pure transform.

Sheet `BOM` (header row 1, then one row per node):
  id          unique key (required)
  parent_id   parent's id; BLANK = the single root FERT
  level       advisory only (for humans / cross-check)
  role        parent | made | bought
  type        FERT | HALB | HAWA | ROH
  name, description(->40 chars), material(existing # or blank=create),
  quantity(def 1), unit(def EA), vendor(-> PIR), price(-> cost), plant(root row only)

Sheet `Operations` (optional; routing, N rows per made node):
  node_id  (a BOM.id: the parent FERT or a made HALB)
  operation, text, work_center, setup_time?, run_time?

Tree is by id/parent_id: exactly ONE root (blank parent_id) = the FERT. Made nodes get nested
`components` + `routing`; bought nodes carry vendor/price -> PIR + cost. genesis builds BOM/routing/PV two
'made' levels deep (FERT -> HALB -> parts); a made node deeper than that is flagged (its own sub-structure
is not built).
"""
import os

_DEF_PLANT = os.getenv("SAP_PLANT", "1710")
_ROLES = {"parent", "made", "bought"}
_TYPES = {"FERT", "HALB", "HAWA", "ROH"}
_MADE_TYPES = {"FERT", "HALB"}


def _s(v):
    return str(v).strip() if v is not None else ""


def _f(v, default=None):
    try:
        return float(v) if _s(v) != "" else default
    except (ValueError, TypeError):
        return default


def _read_sheet(ws):
    """Header row (row 1) -> list of {col_lower: value} dicts; skip fully-blank rows."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_s(h).lower() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(_s(v) == "" for v in r):
            continue
        out.append({h: v for h, v in zip(headers, r) if h})
    return out


def _row_to_node(r):
    """Normalize one BOM row into a node dict (same coercions as genesis_from_csv)."""
    return {
        "id": _s(r.get("id")),
        "parent_id": _s(r.get("parent_id")),
        "role": _s(r.get("role")).lower(),
        "type": _s(r.get("type")).upper() or "HAWA",
        "name": _s(r.get("name")),
        "description": _s(r.get("description")) or _s(r.get("name")),
        "material": _s(r.get("material")) or None,
        "quantity": _f(r.get("quantity"), 1),
        "unit": _s(r.get("unit")) or "EA",
        "vendor": _s(r.get("vendor")) or None,
        "price": _f(r.get("price"), None),
        "plant": _s(r.get("plant")) or None,
    }


def _ops_by_node(op_rows):
    """Group Operations rows by node_id -> {id: [ {operation,text,work_center,setup_time?,run_time?} ]}."""
    ops = {}
    for r in op_rows:
        nid = _s(r.get("node_id"))
        if not nid:
            continue
        op = {"operation": _s(r.get("operation")), "text": _s(r.get("text")),
              "work_center": _s(r.get("work_center"))}
        st, rt = _f(r.get("setup_time")), _f(r.get("run_time"))
        if st is not None:
            op["setup_time"] = st
        if rt is not None:
            op["run_time"] = rt
        ops.setdefault(nid, []).append(op)
    return ops


def _validate(nodes, ops, warnings):
    """Fatal structure problems raise ValueError; soft issues append to `warnings`. Returns the root node."""
    for n in nodes:
        if not n["id"]:
            raise ValueError("BOM: a row has a blank `id` -- every node needs a unique id.")
    ids = [n["id"] for n in nodes]
    dups = sorted({x for x in ids if ids.count(x) > 1})
    if dups:
        raise ValueError(f"BOM: duplicate id(s): {dups}")
    idset = set(ids)
    roots = [n for n in nodes if not n["parent_id"]]
    if len(roots) != 1:
        raise ValueError(f"BOM: expected exactly ONE root (blank parent_id); found {len(roots)}: "
                         f"{[n['id'] for n in roots]}")
    root = roots[0]
    if root["type"] != "FERT":
        warnings.append(f"root '{root['id']}' type is {root['type']}, expected FERT.")
    for n in nodes:
        if n["parent_id"] and n["parent_id"] not in idset:
            raise ValueError(f"BOM: node '{n['id']}' references a missing parent_id '{n['parent_id']}'.")
    parent_of = {n["id"]: n["parent_id"] for n in nodes}
    for n in nodes:                                      # cycle check
        seen, cur = set(), n["id"]
        while cur:
            if cur in seen:
                raise ValueError(f"BOM: cycle detected at node '{n['id']}'.")
            seen.add(cur)
            cur = parent_of.get(cur) or ""
    for n in nodes:
        if n["role"] and n["role"] not in _ROLES:
            warnings.append(f"node '{n['id']}' role '{n['role']}' not in {sorted(_ROLES)}.")
        if n["type"] not in _TYPES:
            warnings.append(f"node '{n['id']}' type '{n['type']}' not in {sorted(_TYPES)}.")
        if _is_bought(n) and not n["vendor"]:
            warnings.append(f"bought node '{n['id']}' has no vendor -> no PIR/cost will be created.")
    for nid in ops:
        if nid not in idset:
            warnings.append(f"Operations: node_id '{nid}' not found in the BOM sheet.")
    children_of = _children_map(nodes)
    def _depth_warn(node, level):
        for ch in children_of.get(node["id"], []):
            if _is_made(ch) and level >= 2:
                warnings.append(f"node '{ch['id']}' is a made sub-assembly at depth {level + 1}; genesis "
                                f"builds BOM/routing/PV only 2 made-levels deep -- its own sub-structure "
                                f"(BOM/routing/PV) will NOT be built.")
            _depth_warn(ch, level + 1)
    _depth_warn(root, 0)
    return root


def _children_map(nodes):
    m = {}
    for n in nodes:
        m.setdefault(n["parent_id"], []).append(n)
    return m


def _is_made(n):
    return n["role"] == "made" or n["type"] in _MADE_TYPES


def _is_bought(n):
    return n["role"] == "bought" or (n["role"] not in ("parent", "made") and n["type"] not in _MADE_TYPES)


def _to_spec_node(n, children_of, ops):
    """One BOM node + its subtree -> the genesis component/child dict (recursive), matching the shape
    _build_made_subassembly and the top-level component loop expect."""
    kids = children_of.get(n["id"], [])
    made = _is_made(n) or bool(kids)
    node = {"name": n["name"], "description": n["description"], "type": n["type"],
            "role": "made" if made else "bought",
            "quantity": n["quantity"], "unit": n["unit"], "material": n["material"]}
    if n["vendor"]:
        node["vendor"] = n["vendor"]
    if n["price"] is not None:
        node["price"] = n["price"]
    if kids:
        node["components"] = [_to_spec_node(k, children_of, ops) for k in kids]
    if n["id"] in ops:
        node["routing"] = ops[n["id"]]
    return node


def genesis_from_excel(path):
    """Parse a tabbed Excel BOM -> (spec, warnings). `spec` is the multi-level dict run_genesis(spec)
    consumes. Read-only; no SAP calls. Raises ValueError on a broken workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl is required to read Excel BOMs -- add it to pyproject.toml.") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = {s.lower(): s for s in wb.sheetnames}
    if "bom" not in sheets:
        raise ValueError(f"workbook has no `BOM` sheet (found: {wb.sheetnames}).")
    nodes = [_row_to_node(r) for r in _read_sheet(wb[sheets["bom"]])]
    op_rows = _read_sheet(wb[sheets["operations"]]) if "operations" in sheets else []
    if not nodes:
        raise ValueError("BOM sheet has no data rows.")
    ops = _ops_by_node(op_rows)
    warnings = []
    root = _validate(nodes, ops, warnings)
    children_of = _children_map(nodes)
    spec = {
        "parent": {"description": root["description"], "type": root["type"] or "FERT",
                   "plant": root["plant"] or _DEF_PLANT, "material": root["material"]},
        "components": [_to_spec_node(k, children_of, ops) for k in children_of.get(root["id"], [])],
    }
    if root["id"] in ops:
        spec["routing"] = ops[root["id"]]
    return spec, warnings
